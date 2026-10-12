"""research/06-power-thermal.xlsx 템플릿 생성기.

값은 비워 두고 수식만 넣는다. 모르는 값은 빈 칸(변수)으로 남긴다.
사용법: python3 tools/make_power_thermal.py [출력 경로]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT = ROOT / "research" / "06-power-thermal.xlsx"

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
CALC_FILL = PatternFill("solid", fgColor="E8F3E8")

COMPONENTS = [
    "디스플레이",
    "광학부 (능동 소자 시)",
    "컴퓨팅 보드",
    "카메라/ToF",
    "햅틱 어레이",
    "조도 센서",
    "방열 (팬 등)",
    "안전 인터록",
    "기타",
]
FIRST_ROW = 2
LAST_ROW = FIRST_ROW + len(COMPONENTS) - 1


def header(ws, titles, widths):
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "B2"


def assumptions(wb):
    ws = wb.active
    ws.title = "가정"
    header(ws, ["변수", "값", "단위", "출처/근거", "확인 날짜"], [28, 12, 10, 40, 14])
    rows = [
        ("일일 운전 시간", 24, "h"),
        ("전원 정격 마진", None, "비율 (예: 0.3)"),
        ("주변 온도 (시험 조건)", None, "°C"),
        ("주변광 (시험 조건)", 500, "lux"),
    ]
    for i, (name, value, unit) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=name)
        cell = ws.cell(row=i, column=2, value=value)
        cell.fill = INPUT_FILL
        ws.cell(row=i, column=3, value=unit)
    ws.cell(row=2, column=4, value="README 0장 (24시간 연속 운전)")
    ws.cell(row=5, column=4, value="README 0장 (1차 목표 환경)")
    note = ws.cell(row=8, column=1, value="노란 칸 = 입력, 초록 칸 = 계산. 빈 입력은 미확인 변수로 취급한다.")
    note.font = Font(italic=True)


def power_budget(wb):
    ws = wb.create_sheet("전력예산")
    titles = [
        "부품", "전원 레일", "전압 V", "평균 전류 A", "피크 전류 A",
        "평균 전력 W", "피크 전력 W", "듀티 (0–1)", "동작 시간 h/일",
        "유효 평균 W", "일일 에너지 Wh", "출처 (데이터시트/실측)", "확인 날짜", "비고",
    ]
    header(ws, titles, [22, 12, 9, 12, 12, 12, 12, 11, 13, 12, 14, 26, 12, 24])

    for r, name in enumerate(COMPONENTS, start=FIRST_ROW):
        ws.cell(row=r, column=1, value=name)
        for col in (2, 3, 4, 5, 8, 12, 13, 14):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        # W = V × A. W만 아는 경우 F/G 칸에 직접 입력해 덮어쓴다.
        ws.cell(row=r, column=6, value=f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})')
        ws.cell(row=r, column=7, value=f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})')
        ws.cell(row=r, column=9, value="=가정!$B$2").fill = INPUT_FILL
        ws.cell(row=r, column=10, value=f'=IF(OR(F{r}="",H{r}=""),"",F{r}*H{r})')
        ws.cell(row=r, column=11, value=f'=IF(OR(J{r}="",I{r}=""),"",J{r}*I{r})')
        for col in (6, 7, 10, 11):
            ws.cell(row=r, column=col).fill = CALC_FILL

    t = LAST_ROW + 2
    summary = [
        ("합계 유효 평균 전력 W", f"=SUM(J{FIRST_ROW}:J{LAST_ROW})"),
        ("합계 피크 전력 W (동시 최악)", f"=SUM(G{FIRST_ROW}:G{LAST_ROW})"),
        ("일일 에너지 Wh", f"=SUM(K{FIRST_ROW}:K{LAST_ROW})"),
        ("일일 에너지 kWh", f"=B{t + 2}/1000"),
        ("필요 전원 정격 W (피크 × (1+마진))",
         f'=IF(가정!B3="","마진 미입력",B{t + 1}*(1+가정!B3))'),
        ("미입력 부품 수 (평균 W 없음)",
         f'=COUNTBLANK(F{FIRST_ROW}:F{LAST_ROW})'),
    ]
    for i, (label, formula) in enumerate(summary):
        ws.cell(row=t + i, column=1, value=label).font = HEADER_FONT
        ws.cell(row=t + i, column=2, value=formula).fill = CALC_FILL

    duty = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                          allow_blank=True, showErrorMessage=True,
                          error="듀티는 0과 1 사이 값이어야 합니다.")
    ws.add_data_validation(duty)
    duty.add(f"H{FIRST_ROW}:H{LAST_ROW}")


def thermal(wb):
    ws = wb.create_sheet("열")
    titles = [
        "측정 위치", "센서 종류", "허용 온도 °C", "쓰로틀 기준 °C", "차단 기준 °C",
        "기준 출처", "24h 시험 최대 °C", "여유 °C", "판정", "비고",
    ]
    header(ws, titles, [24, 14, 12, 13, 12, 26, 15, 10, 10, 24])
    points = [
        "컴퓨팅 보드 SoC",
        "디스플레이 패널/백라이트",
        "햅틱 어레이 구동부",
        "전원 공급 장치",
        "외함 표면 (사용자 접촉부)",
        "흡기",
        "배기",
    ]
    for r, name in enumerate(points, start=2):
        ws.cell(row=r, column=1, value=name)
        for col in (2, 3, 4, 5, 6, 7, 10):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.cell(row=r, column=8, value=f'=IF(OR(C{r}="",G{r}=""),"",C{r}-G{r})').fill = CALC_FILL
        ws.cell(row=r, column=9,
                value=f'=IF(H{r}="","미측정",IF(H{r}>=0,"합격","불합격"))').fill = CALC_FILL
    ws.cell(row=len(points) + 3, column=1,
            value="허용·차단 기준은 데이터시트 또는 해당 안전 기준 출처를 반드시 기록한다.").font = Font(italic=True)


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    wb = Workbook()
    assumptions(wb)
    power_budget(wb)
    thermal(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
